import zipfile
import xml.etree.ElementTree as ET

import docx
import pytest

from ligaotai.export import export_book, export_path
from ligaotai.fsutil import write_json
from ligaotai.scenes import BrokenSceneFile

SK = {"generated": "x", "by": "author", "volumes": [{"title": "第一卷 起", "chapters": [
    {"title": "开篇", "notes": [], "items": [
        {"type": "scene", "id": "S-0001", "thread": "L-001"},
        {"type": "hole", "id": "H-001", "task": "在 S-0001 与 S-0003 之间补写：\n大闹天宫"},
        {"type": "scene", "id": "S-0099", "thread": "L-001"}]}]}],
    "unplaced": {"scenes": [{"id": "S-0005", "thread": "L-002", "why": "no_time"}], "holes": []}}

MD = """# 第一卷 起

## 开篇

<!-- S-0001 -->
S-0001 的正文。

> 【空洞 H-001】在 S-0001 与 S-0003 之间补写： 大闹天宫

> 【缺失场景 S-0099：原稿里已经没有这一块了】

# 附：未定位

<!-- S-0005 -->
S-0005 的正文。
"""

TXT = """第一卷 起

开篇

S-0001 的正文。

【空洞 H-001】在 S-0001 与 S-0003 之间补写： 大闹天宫

【缺失场景 S-0099：原稿里已经没有这一块了】

附：未定位

S-0005 的正文。
"""


def test_导出md和txt(book_with_threads):
    b = book_with_threads
    write_json(b.skeleton_path, SK)
    r = export_book(b)
    assert r == {"md": "导出/测试书.md", "txt": "导出/测试书.txt", "docx": "导出/测试书.docx",
                 "epub": "导出/测试书.epub", "scenes": 2, "holes": 1, "missing": 1, "chars": 22, "cut": 0}
    assert export_path(b, "md").read_text(encoding="utf-8") == MD
    assert export_path(b, "txt").read_text(encoding="utf-8") == TXT


def test_没有骨架不能导出(book_with_threads):
    with pytest.raises(FileNotFoundError):
        export_book(book_with_threads)


def test_格式只认四种(book_with_threads):
    for fmt in ("md", "txt", "docx", "epub"):
        export_path(book_with_threads, fmt)
    with pytest.raises(ValueError):
        export_path(book_with_threads, "pdf")


def test_导出docx_卷章是标题_空洞和缺失照样标出来(book_with_threads):
    b = book_with_threads
    write_json(b.skeleton_path, SK)
    export_book(b)
    d = docx.Document(str(export_path(b, "docx")))
    paras = [(p.style.name, p.text) for p in d.paragraphs if p.text]
    assert paras == [
        ("Title", "测试书"),
        ("Heading 1", "第一卷 起"),
        ("Heading 2", "开篇"),
        ("Normal", "S-0001 的正文。"),
        ("Normal", "【空洞 H-001】在 S-0001 与 S-0003 之间补写： 大闹天宫"),
        ("Normal", "【缺失场景 S-0099：原稿里已经没有这一块了】"),
        ("Heading 1", "附：未定位"),
        ("Normal", "S-0005 的正文。"),
    ]
    assert d.core_properties.title == "测试书"


def test_导出docx_场景里的换行分成段落(book_with_threads):
    b = book_with_threads
    (b.scenes_dir / "S-0001.md").write_text(
        (b.scenes_dir / "S-0001.md").read_text(encoding="utf-8").replace("S-0001 的正文。", "第一段。" + chr(10) * 2 + "第二段。" + chr(10) + "第三段。"),
        encoding="utf-8")
    write_json(b.skeleton_path, SK)
    export_book(b)
    texts = [p.text for p in docx.Document(str(export_path(b, "docx"))).paragraphs]
    assert texts[texts.index("第一段。"):texts.index("第一段。") + 3] == ["第一段。", "第二段。", "第三段。"]


def test_分段_硬换行的正文把折行接回去():
    """验收书原文是固定宽度硬换行的，一行当一段会把一句话拆成两段（9-28 真书截图看到）。
    贴满宽度的行跟下一行接上，没贴满的才是段落结尾；标题行短，自成一段。"""
    from ligaotai.export import _paras
    w = "甲" * 20
    text = chr(10).join(["第四回 标题", w, w, "乙乙乙。", w, "丙丙。」"])
    assert _paras(text) == ["第四回 标题", w + w + "乙乙乙。", w + "丙丙。」"]


def test_分段_不是硬换行的正文一行一段():
    from ligaotai.export import _paras
    text = chr(10).join(["短的一段。", "另一段长一点点的内容。", "", "第三段。"])
    assert _paras(text) == ["短的一段。", "另一段长一点点的内容。", "第三段。"]


def test_分段_段首有缩进的稿子按缩进分段():
    """雪月梅那种长短行交替的折行，宽度判断不管用；但段首都有「　　」缩进，按缩进分最稳。"""
    from ligaotai.export import _paras
    ind = chr(0x3000) * 2
    text = chr(10).join(["第一回 标题", "", ind + "卻說為人在世，所以這天、", "地、君、親、師的大恩，", "心，思所報答。",
                         ind + "第二段開頭，", "接著寫完。"])
    assert _paras(text) == ["第一回 标题", "卻說為人在世，所以這天、地、君、親、師的大恩，心，思所報答。", "第二段開頭，接著寫完。"]


def test_分段_以句读标点开头的行隔着空行也接回上一段():
    from ligaotai.export import _paras
    text = chr(10).join(["他說道：「好", "", "。」眾人都笑了。"])
    assert _paras(text) == ["他說道：「好。」眾人都笑了。"]


def test_分段_英文折行接回去时补空格():
    from ligaotai.export import _paras
    a = "a" * 39 + "b"
    text = chr(10).join([a, a, "end."])
    assert _paras(text) == [a + " " + a + " end."]


def _epub(b):
    z = zipfile.ZipFile(export_path(b, "epub"))
    return z, {n: z.read(n).decode("utf-8") for n in z.namelist()}


def test_导出epub_结构合规(book_with_threads):
    b = book_with_threads
    write_json(b.skeleton_path, SK)
    export_book(b)
    z, files = _epub(b)
    first = z.infolist()[0]
    # EPUB 规定：mimetype 是第一个文件、不压缩、内容就是这一串
    assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
    assert files["mimetype"] == "application/epub+zip"
    assert "OEBPS/content.opf" in files["META-INF/container.xml"]
    # 所有 xml / xhtml 都得是合法 XML
    for n, t in files.items():
        if n.endswith((".xml", ".opf", ".xhtml", ".ncx")):
            ET.fromstring(t.encode("utf-8"))
    opf = ET.fromstring(files["OEBPS/content.opf"].encode("utf-8"))
    ns = {"o": "http://www.idpf.org/2007/opf", "dc": "http://purl.org/dc/elements/1.1/"}
    assert opf.find(".//dc:title", ns).text == "测试书"
    assert opf.find(".//dc:language", ns).text == "zh-CN"
    hrefs = {i.get("id"): i.get("href") for i in opf.findall(".//o:manifest/o:item", ns)}
    spine = [hrefs[r.get("idref")] for r in opf.findall(".//o:spine/o:itemref", ns)]
    for h in hrefs.values():  # 清单里列的文件都真的在包里
        assert f"OEBPS/{h}" in files
    body = "".join(files[f"OEBPS/{h}"] for h in spine)
    order = ["第一卷 起", "开篇", "S-0001 的正文。", "【空洞 H-001】", "【缺失场景 S-0099", "附：未定位", "S-0005 的正文。"]
    at = [body.index(x) for x in order]
    assert at == sorted(at)  # 阅读顺序跟骨架一致
    nav = files["OEBPS/nav.xhtml"]
    assert "第一卷 起" in nav and "开篇" in nav and "附：未定位" in nav


def test_导出epub_正文里的尖括号和与号要转义(book_with_threads):
    b = book_with_threads
    (b.scenes_dir / "S-0001.md").write_text(
        (b.scenes_dir / "S-0001.md").read_text(encoding="utf-8").replace("S-0001 的正文。", "他说<好>&走"),
        encoding="utf-8")
    write_json(b.skeleton_path, SK)
    export_book(b)
    _, files = _epub(b)
    body = "".join(t for n, t in files.items() if n.endswith(".xhtml"))
    assert "他说&lt;好&gt;&amp;走" in body


def test_场景文件坏了_导出往上抛不当成原稿删了(book_with_threads):
    """S2：export.py:322 原先吞了 BrokenSceneFile，把「文件手改坏了」说成「原稿里已经
    没有这一块了」——两件不一样的事，得让它往上抛，接口那层才能报 500 说清是哪个文件坏的。"""
    b = book_with_threads
    write_json(b.skeleton_path, {"volumes": [{"title": "卷", "chapters": [{"title": "章", "items": [
        {"type": "scene", "id": "S-0001"}]}]}]})
    (b.scenes_dir / "S-0001.md").write_text("没有头信息，被手改坏了", encoding="utf-8")
    with pytest.raises(BrokenSceneFile):
        export_book(b)
