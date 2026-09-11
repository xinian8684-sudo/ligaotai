from docx import Document

from ligaotai.readers import SUPPORTED, detect_encoding, read_text

TEXT = "林清年方十六，住在青州城外。那一年雪下得很大。\n第二段。"


def test_supported():
    assert SUPPORTED == {".txt", ".md", ".docx"}


def test_utf8(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes(TEXT.encode("utf-8"))
    assert read_text(p) == (TEXT, "utf-8")


def test_utf8_bom_removed(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes(TEXT.encode("utf-8-sig"))
    text, enc = read_text(p)
    assert enc == "utf-8-sig"
    assert text == TEXT


def test_gbk(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes((TEXT * 3).encode("gbk"))
    assert read_text(p) == (TEXT * 3, "gb18030")


def test_utf16_bom(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes(TEXT.encode("utf-16"))
    text, enc = read_text(p)
    assert enc == "utf-16"
    assert text == TEXT


def test_crlf_normalized(tmp_path):
    p = tmp_path / "a.md"
    p.write_bytes("第一行\r\n第二行\r第三行".encode("utf-8"))
    assert read_text(p)[0] == "第一行\n第二行\n第三行"


def test_ascii_is_utf8():
    assert detect_encoding(b"hello") == "utf-8"


def test_docx_headings_become_markdown(tmp_path):
    doc = Document()
    doc.add_heading("全书名", level=0)
    doc.add_heading("第一章 开端", level=1)
    doc.add_paragraph("正文一")
    doc.add_paragraph("")
    doc.add_heading("小节", level=2)
    doc.add_paragraph("正文二")
    p = tmp_path / "a.docx"
    doc.save(str(p))
    text, enc = read_text(p)
    assert enc == "docx"
    assert text == "# 全书名\n# 第一章 开端\n正文一\n\n## 小节\n正文二"
