import pytest
from docx import Document

from ligaotai.book import create_book
from ligaotai.fsutil import read_json
from ligaotai.importer import run_import


def make_src(tmp_path):
    src = tmp_path / "我的稿子"
    (src / "旧稿").mkdir(parents=True)
    (src / ".obsidian").mkdir()
    (src / "a.txt").write_bytes("第一章 开端\n林清年方十六。".encode("gbk"))
    (src / "旧稿" / "b.md").write_text("# 第二章\n雪夜。", encoding="utf-8")
    doc = Document()
    doc.add_heading("第三章", level=1)
    doc.add_paragraph("正文。")
    doc.save(str(src / "c.docx"))
    (src / "d.pdf").write_bytes(b"%PDF")
    (src / ".obsidian" / "e.txt").write_text("x", encoding="utf-8")
    (src / "~$f.docx").write_bytes(b"lock")
    return src


def test_first_import(tmp_path, book):
    src = make_src(tmp_path)
    calls = []
    summary = run_import(book, src, lambda done, total, *a: calls.append((done, total)))
    assert summary["added"] == 3
    assert summary["changed"] == 0
    assert summary["failed"] == []
    assert sorted(s["path"] for s in summary["skipped"]) == [".obsidian/e.txt", "d.pdf", "~$f.docx"]
    assert (book.originals_dir / "我的稿子" / "旧稿" / "b.md").exists()
    files = read_json(book.manifest_path)["files"]
    assert set(files) == {"我的稿子/a.txt", "我的稿子/旧稿/b.md", "我的稿子/c.docx"}
    assert files["我的稿子/a.txt"]["encoding"] == "gb18030"
    assert files["我的稿子/c.docx"]["encoding"] == "docx"
    assert book.step("import")["status"] == "done"
    assert calls[-1] == (6, 6)


def test_reimport_unchanged_keeps_downstream(tmp_path, book):
    src = make_src(tmp_path)
    run_import(book, src)
    book.set_step("split", "done")
    summary = run_import(book, src)
    assert (summary["added"], summary["changed"], summary["unchanged"]) == (0, 0, 3)
    assert book.step("split")["status"] == "done"


def test_reimport_changed_outdates_downstream(tmp_path, book):
    src = make_src(tmp_path)
    run_import(book, src)
    book.set_step("split", "done")
    (src / "旧稿" / "b.md").write_text("# 第二章\n雪夜，改过了。", encoding="utf-8")
    summary = run_import(book, src)
    assert summary["changed"] == 1
    assert book.step("split")["status"] == "outdated"
    copied = (book.originals_dir / "我的稿子" / "旧稿" / "b.md").read_text(encoding="utf-8")
    assert "改过了" in copied


def test_broken_docx_is_reported_not_copied(tmp_path, book):
    src = tmp_path / "稿"
    src.mkdir()
    (src / "坏.docx").write_bytes(b"not a zip")
    summary = run_import(book, src)
    assert [f["path"] for f in summary["failed"]] == ["稿/坏.docx"]
    assert not (book.originals_dir / "稿" / "坏.docx").exists()
    assert read_json(book.manifest_path)["files"] == {}


def test_not_a_folder(tmp_path, book):
    with pytest.raises(NotADirectoryError):
        run_import(book, tmp_path / "没有")


def test_skips_nested_book_library(tmp_path, book):
    src = tmp_path / "大文件夹"
    src.mkdir()
    (src / "a.txt").write_text("正文。", encoding="utf-8")
    nested = src / "别的书库" / "某本书"
    nested.mkdir(parents=True)
    (nested / "book.json").write_text("{}", encoding="utf-8")
    (nested / "笔记.md").write_text("笔记内容。", encoding="utf-8")
    summary = run_import(book, src)
    assert summary["added"] == 1
    files = read_json(book.manifest_path)["files"]
    assert set(files) == {"大文件夹/a.txt"}
    reasons = {s["path"]: s["reason"] for s in summary["skipped"]}
    assert reasons["别的书库/某本书/book.json"] == "理稿台书库"
    assert reasons["别的书库/某本书/笔记.md"] == "理稿台书库"


def test_import_own_folder_raises(book):
    book.originals_dir.mkdir(parents=True)
    with pytest.raises(ValueError):
        run_import(book, book.root / "原稿")


def test_reimport_restores_deleted_copy(tmp_path, book):
    src = make_src(tmp_path)
    run_import(book, src)
    copy = book.originals_dir / "我的稿子" / "a.txt"
    assert copy.exists()
    copy.unlink()
    summary = run_import(book, src)
    assert copy.exists()
    assert summary["changed"] == 1


def test_manifest记下每个原稿根目录的来源路径(tmp_path):
    """同名但不同位置的两个文件夹会共用一个 root_name，
    manifest 要记住它上次是从哪来的，界面才提醒得了。"""
    lib = tmp_path / "书库"
    lib.mkdir()
    b = create_book(lib, "测试书")

    src = tmp_path / "甲" / "我的稿子"
    src.mkdir(parents=True)
    (src / "a.txt").write_text("正文", encoding="utf-8")

    run_import(b, src)
    manifest = read_json(b.manifest_path, {})
    assert manifest["roots"]["我的稿子"] == str(src.resolve())


def test_同名不同路径再导入时来源路径会被更新(tmp_path):
    lib = tmp_path / "书库"
    lib.mkdir()
    b = create_book(lib, "测试书")
    甲 = tmp_path / "甲" / "我的稿子"
    乙 = tmp_path / "乙" / "我的稿子"
    for d in (甲, 乙):
        d.mkdir(parents=True)
        (d / "a.txt").write_text("正文", encoding="utf-8")

    run_import(b, 甲)
    run_import(b, 乙)
    manifest = read_json(b.manifest_path, {})
    assert manifest["roots"]["我的稿子"] == str(乙.resolve())
