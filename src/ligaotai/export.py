"""按骨架拼书，导出 md / txt / docx / epub（计划④ spec 第 8 节）。只导出，不回写原稿。"""

from __future__ import annotations

import uuid
import zipfile
from datetime import datetime, timezone
from html import escape
from pathlib import Path

import docx
from docx.shared import RGBColor

from .book import Book
from .fsutil import atomic_write_text, safe_name
from .scenes import BrokenSceneFile, get_scene
from .skeleton import load_skeleton
from .threads_ops import BrokenThreadsFile, load_threads
from .triage import BrokenBoardFile, columns, version_map

FORMATS = ("md", "txt", "docx", "epub")


def export_path(book: Book, fmt: str) -> Path:
    if fmt not in FORMATS:
        raise ValueError("只能导出 md、txt、docx 或 epub")
    title = book.load().get("title") or book.name
    return book.export_dir / f"{safe_name(title)}.{fmt}"


def _text(book: Book, sid: str) -> str | None:
    try:
        sc = get_scene(book, sid)
    except BrokenSceneFile:
        # BrokenSceneFile 继承 ValueError，必须排在下面那条前面，不然被吃成「原稿里已经
        # 没有这一块了」——场景文件手改坏了跟场景真的被删掉，是两件不一样的事，前者要
        # 让接口报 500、把哪个文件坏了说清楚，不能悄悄说成后者（S2）。
        raise
    except (ValueError, FileNotFoundError):
        return None
    return None if sc.removed else sc.text.strip()


def _cut_lookup(book: Book) -> tuple[dict, dict]:
    """场景所属线现在是不是被砍掉了——只用来给导出结果附带一个 cut 计数提醒界面，
    线 / 看板文件不存在或者坏了就当查不出来（不影响导出本身，导出不依赖归线结果）。"""
    try:
        threads = load_threads(book)
        cols = columns(book, threads)
    except (FileNotFoundError, BrokenThreadsFile, BrokenBoardFile):
        return {}, {}
    vmap = version_map(book)
    thread_of: dict[str, str] = {}
    for t in threads.get("threads") or []:
        if isinstance(t, dict) and t.get("id"):
            for sid0 in t.get("scenes") or []:
                thread_of.setdefault(sid0, t["id"])
                thread_of.setdefault(vmap.get(sid0, sid0), t["id"])
    return cols, thread_of


def export_book(book: Book) -> dict:
    sk = load_skeleton(book)
    vmap = version_map(book)
    cols, thread_of = _cut_lookup(book)
    md: list[str] = []
    txt: list[str] = []
    # docx / epub 用的结构：("volume", 标题) / ("chapter", 标题) / ("scene", 正文) / ("note", 提示)
    # note 是空洞和缺失场景的提示，不是原文
    nodes: list[tuple[str, str]] = []
    counts = {"scenes": 0, "holes": 0, "missing": 0, "chars": 0, "cut": 0}

    def emit(it: dict) -> None:
        if it.get("type") == "hole":
            counts["holes"] += 1
            task = " ".join(str(it.get("task") or "").splitlines())
            md.append(f"> 【空洞 {it['id']}】{task}")
            txt.append(f"【空洞 {it['id']}】{task}")
            nodes.append(("note", txt[-1]))
            return
        sid = it.get("id")
        # M3：场景正文取组里现在的主版本，不是骨架里原样存的那个引用——骨架可能是在作者
        # 换主版本之前生成的，spec 8 要求导出的是主版本原文。
        real = vmap.get(sid, sid)
        text = _text(book, real)
        if text is None:
            counts["missing"] += 1
            md.append(f"> 【缺失场景 {sid}：原稿里已经没有这一块了】")
            txt.append(f"【缺失场景 {sid}：原稿里已经没有这一块了】")
            nodes.append(("note", txt[-1]))
            return
        counts["scenes"] += 1
        counts["chars"] += len(text)
        if cols.get(thread_of.get(sid), {}).get("col") == "cut":
            counts["cut"] += 1
        md.append(f"<!-- {real} -->\n{text}")
        txt.append(text)
        nodes.append(("scene", text))

    for v in sk.get("volumes") or []:
        md.append(f"# {v['title']}")
        txt.append(v["title"])
        nodes.append(("volume", v["title"]))
        for ch in v.get("chapters") or []:
            md.append(f"## {ch['title']}")
            txt.append(ch["title"])
            nodes.append(("chapter", ch["title"]))
            for it in ch.get("items") or []:
                emit(it)
    up = sk.get("unplaced") or {}
    rest = [{"type": "scene", "id": x.get("id")} for x in up.get("scenes") or []] + \
           [{**x, "type": "hole"} for x in up.get("holes") or []]
    if rest:
        md.append("# 附：未定位")
        txt.append("附：未定位")
        nodes.append(("volume", "附：未定位"))
        for it in rest:
            emit(it)

    md_path, txt_path = export_path(book, "md"), export_path(book, "txt")
    atomic_write_text(md_path, "\n\n".join(md) + "\n")
    atomic_write_text(txt_path, "\n\n".join(txt) + "\n")
    title = book.load().get("title") or book.name
    docx_path, epub_path = export_path(book, "docx"), export_path(book, "epub")
    _write_docx(docx_path, title, nodes)
    _write_epub(epub_path, title, nodes)
    rel = lambda p: p.relative_to(book.root).as_posix()  # noqa: E731
    return {"md": rel(md_path), "txt": rel(txt_path), "docx": rel(docx_path), "epub": rel(epub_path), **counts}


_SENT_END = set("。！？」』”…）)!?.：:；;﹔—")
_NO_START = set("。，、；：」』！？）﹔")


def _paras(text: str) -> list[str]:
    """场景正文拆成段落（docx / epub 用；md / txt 照原文不动）。空行永远是段落分隔。

    稿子常是硬换行的（验收书就是，9-28 截图看到一句话被拆成两段），但折法不一样：西游记按固定
    宽度折，雪月梅长短行交替（44 字、20 字轮着来）、段首有「　　」缩进，西游记的诗词中间也有
    「　　」。猜哪条规则适用不靠谱，于是三种分法都算一遍，挑毛病最少的：
    - 一行一段；
    - 按宽度接：贴满最长行宽度的行跟下一行接上；
    - 按缩进接：缩进的行开新段，不缩进的行接到上一段。
    毛病＝句子被拆开（段尾不是句末标点）＋在句末标点后面硬接（把两段并成了一段）。
    只数前一种的话「全书并成一段」是满分，所以两头都数。"""
    raw = text.splitlines()
    lines = [ln.strip() for ln in raw]
    lens = sorted(len(ln) for ln in lines if ln)
    if not lens:
        return []
    indented = [bool(ln) and r[:2] in ("　　", "  ") for ln, r in zip(lines, raw)]
    width = lens[-1]

    def build(joins) -> tuple[int, list[str]]:
        out: list[str] = []
        bad = 0
        joining = False
        for i, ln in enumerate(lines):
            if not ln:
                joining = False
                continue
            if out and ln[0] in _NO_START:  # 中文不会拿「。」「，」开一段：隔着空行也接回去
                out[-1] += ln
            elif joining and joins(i, ln):
                prev = out[-1]
                bad += prev[-1] in _SENT_END
                gap = " " if prev[-1].isascii() and prev[-1].isalnum() and ln[0].isascii() and ln[0].isalnum() else ""
                out[-1] = prev + gap + ln
            else:
                out.append(ln)
            joining = True
        bad += sum(p[-1] not in _SENT_END for p in out[:-1])
        return bad, out

    options = [build(lambda i, ln: False)]
    if width >= 16:  # 再短的「宽度」多半只是几行短句
        options.append(build(lambda i, ln: len(lines[_prev(lines, i)]) >= width * 0.9))
    if sum(indented) >= 2:
        options.append(build(lambda i, ln: not indented[i]))
    return min(options, key=lambda o: o[0])[1]  # 平手取靠前的（越靠前越保守）


def _prev(lines: list[str], i: int) -> int:
    """i 前面最近一个非空行的下标（build 只在「前面接着有行」时才问，一定找得到）。"""
    j = i - 1
    while not lines[j]:
        j -= 1
    return j


def _write_docx(path: Path, title: str, nodes: list[tuple[str, str]]) -> None:
    d = docx.Document()
    d.core_properties.title = title
    d.add_heading(title, level=0)
    prev = None
    for kind, val in nodes:
        if kind == "volume":
            d.add_heading(val, level=1)
        elif kind == "chapter":
            d.add_heading(val, level=2)
        elif kind == "scene":
            if prev == "scene":
                d.add_paragraph()  # 场景之间空一行
            for para in _paras(val):
                d.add_paragraph(para)
        else:
            run = d.add_paragraph().add_run(val)
            run.italic = True
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        prev = kind
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    d.save(str(tmp))
    tmp.replace(path)  # 先写临时文件再换过去，写一半崩了不留半个文件


_XHTML = """<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="zh-CN" lang="zh-CN">
<head><meta charset="utf-8"/><title>{title}</title><link rel="stylesheet" type="text/css" href="style.css"/></head>
<body>
{body}
</body>
</html>
"""
_CSS = ("body{line-height:1.8}p{text-indent:2em;margin:0}"
        "p.note{text-indent:0;color:#888;font-style:italic;margin:.8em 0}p.gap{height:1em}h1,h2{text-align:center}")


def _write_epub(path: Path, title: str, nodes: list[tuple[str, str]]) -> None:
    """EPUB 3（另带一份 NCX 照顾老阅读器）。一卷一个开篇页、一章一个文件。自己拼 zip，
    不引 ebooklib——它是 AGPL，跟本仓库的 MIT 不搭。"""
    pages: list[dict] = []

    def page(t: str, level: int) -> dict:
        p = {"file": f"p{len(pages) + 1:04d}.xhtml", "title": t, "level": level,
             "body": [f"<h{level}>{escape(t)}</h{level}>"], "last": None}
        pages.append(p)
        return p

    cur = None
    for kind, val in nodes:
        if kind == "volume":
            cur = page(val, 1)
        elif kind == "chapter":
            cur = page(val, 2)
        else:
            if cur is None:  # 骨架第一卷之前不会有正文，防御一下
                cur = page(title, 1)
            if kind == "scene":
                if cur["last"] == "scene":
                    cur["body"].append('<p class="gap"></p>')
                cur["body"] += [f"<p>{escape(x)}</p>" for x in _paras(val)]
            else:
                cur["body"].append(f'<p class="note">{escape(val)}</p>')
            cur["last"] = kind
    if not pages:
        page(title, 1)

    # 目录：章挂在它前面那一卷下面
    nav: list[str] = []
    in_vol = False
    for p in pages:
        link = f'<a href="{p["file"]}">{escape(p["title"])}</a>'
        if p["level"] == 1:
            if in_vol:
                nav.append("</ol></li>")
            nav.append(f"<li>{link}<ol>")
            in_vol = True
        else:
            nav.append(f"<li>{link}</li>")
    if in_vol:
        nav.append("</ol></li>")
    nav_body = ('<nav epub:type="toc" id="toc"><h1>目录</h1><ol>'
                + "".join(nav).replace("<ol></ol>", "") + "</ol></nav>")
    ncx_points = "".join(
        f'<navPoint id="n{i}" playOrder="{i}"><navLabel><text>{escape(p["title"])}</text></navLabel>'
        f'<content src="{p["file"]}"/></navPoint>' for i, p in enumerate(pages, 1))

    uid = f"urn:uuid:{uuid.uuid5(uuid.NAMESPACE_URL, 'ligaotai:' + title)}"
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    manifest = "".join(f'<item id="i{i}" href="{p["file"]}" media-type="application/xhtml+xml"/>'
                       for i, p in enumerate(pages, 1))
    spine = "".join(f'<itemref idref="i{i}"/>' for i in range(1, len(pages) + 1))
    opf = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid" xml:lang="zh-CN">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
<dc:identifier id="uid">{uid}</dc:identifier><dc:title>{escape(title)}</dc:title><dc:language>zh-CN</dc:language>
<meta property="dcterms:modified">{now}</meta>
</metadata>
<manifest>
<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>
<item id="css" href="style.css" media-type="text/css"/>
{manifest}
</manifest>
<spine toc="ncx">{spine}</spine>
</package>
"""
    ncx = f"""<?xml version="1.0" encoding="utf-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
<head><meta name="dtb:uid" content="{uid}"/></head>
<docTitle><text>{escape(title)}</text></docTitle>
<navMap>{ncx_points}</navMap>
</ncx>
"""
    container = """<?xml version="1.0" encoding="utf-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        # EPUB 规定 mimetype 必须是第一个文件、不压缩
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", container)
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/toc.ncx", ncx)
        z.writestr("OEBPS/style.css", _CSS)
        z.writestr("OEBPS/nav.xhtml", _XHTML.format(title="目录", body=nav_body))
        for p in pages:
            z.writestr(f"OEBPS/{p['file']}", _XHTML.format(title=escape(p["title"]), body="\n".join(p["body"])))
    tmp.replace(path)
